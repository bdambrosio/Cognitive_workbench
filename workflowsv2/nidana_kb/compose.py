#!/usr/bin/env python3
"""Write a translated chapter in the translation format (TRANSLATION_FORMAT.md)
from the team's Sanskrit edition of the chapter and the English written for it.

    python3 workflowsv2/nidana_kb/compose.py <edition.docx> <translation dir> <out.docx> \\
        --chapter 1 --title-en "The five means of diagnosis"

THE EDITION is the team's file for one chapter: a Heading 1 title, then per
section a Heading 2, the root verses under मूलम्, the Madhukośa under
मधुकोशव्याख्या and the Ātaṅkadarpaṇa under आतङ्कदर्पणव्याख्या, each in
Devanagari followed by its IAST. `read_edition` reads it.

THE TRANSLATION DIRECTORY holds one `sec_NN.json` per section (NN its index
in the edition): heading_en, translation, padaccheda, words, notes, and the
commentaries either as `madhukosa` / `atanka`, lists of {k, en} answering the
numbered segments of the work file (--work-file), or, in the first two
chapters, as whole texts `madhukosa_en` / `atanka_en`. Segments are the
back-index: each English segment is written beside its Sanskrit with the
same [k]. A section with no file is written with its Sanskrit only
and a `Note:` saying it is untranslated.

The Sanskrit is copied from the edition, never retyped: the verse numbers
`||n||` become `॥ n ॥`, and each verse is its own `Verse:` and `IAST:`
paragraph. A verse the edition splits across two sections (the first half
ends |n|) is numbered na and nb. A source reference after a verse number
becomes the `Source:` line. The colophon section is written as a `Note:`.
"""
from __future__ import annotations

import argparse
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Dict, List, Tuple

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
PARTS = {"मूलम्": "mula", "मधुकोशव्याख्या": "madhukosa", "आतङ्कदर्पणव्याख्या": "atanka"}
DEVA_DIGITS = "०१२३४५६७८९"
#: A verse number: ||n|| (or ॥ n ॥) ends a verse; |n| ends the first half of
#: a verse whose second half, ||n||, opens the next section.
END_RE = re.compile(r"(\|\||॥|\|)\s*([०-९0-9]+)\s*(?:\|\||॥|\|)")
#: A source reference right after a verse number: (वा. नि. अ. १) or (vā. ni. a. 1).
#: It is abbreviated, so it holds a full stop; a parenthesised line of verse
#: text (the edition has some) does not, and stays with the verse text.
SOURCE_RE = re.compile(r"\s*\(([^().]{1,20}\.[^()]{0,30})\)\s*\|?")
#: The edition's heading for the chapter's colophon.
COLOPHON = "पुष्पिका"


def _deva(t: str) -> bool:
    """Whether a paragraph is Devanagari rather than its IAST twin: most of
    its letters are Devanagari. An IAST paragraph can carry a stray danda or
    Devanagari numeral, so any Devanagari at all is not the test."""
    letters = [c for c in t if c.isalpha()]
    return bool(letters) and sum(1 for c in letters if "ऀ" <= c <= "ॿ") > len(letters) / 2


def read_edition(path: Path) -> Dict[str, Any]:
    root = ET.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    out: Dict[str, Any] = {"file": Path(path).name, "title_sa": "", "title_iast": "", "sections": []}
    sec, part, after_title = None, None, False
    for p in root.find(W + "body").iter(W + "p"):
        ppr = p.find(W + "pPr")
        style = ""
        if ppr is not None and ppr.find(W + "pStyle") is not None:
            style = ppr.find(W + "pStyle").get(W + "val") or ""
        t = "".join(x.text or "" for x in p.iter(W + "t")).strip()
        if not t:
            continue
        if style.startswith("Heading1"):
            out["title_sa"], after_title = t, True
            continue
        if after_title and not _deva(t):
            out["title_iast"], after_title = t, False
            continue
        if style.startswith("Heading2"):
            sec = {"head_sa": t, "head_iast": "", **{k: {"sa": "", "iast": ""} for k in PARTS.values()}}
            out["sections"].append(sec)
            part = "head"
            continue
        if t in PARTS:
            part = PARTS[t]
            continue
        if sec is None:
            continue
        if part == "head":
            sec["head_iast"] = t
            continue
        key = "sa" if _deva(t) else "iast"
        sec[part][key] = (sec[part][key] + "\n" + t).strip()
    return out


def split_verses(text: str, halves: Dict[str, str]) -> List[Dict[str, str]]:
    """The verses of a passage: {"n", "text", "source"}, n in Arabic digits.
    A first half (|n|) is numbered na and the second half that follows it
    nb; `halves` carries the halves already seen across the chapter's
    sections. Text after the last number is joined to the verse before it."""
    out: List[Dict[str, str]] = []
    start = 0
    for m in END_RE.finditer(text):
        n = m.group(2).translate(str.maketrans(DEVA_DIGITS, "0123456789"))
        if m.group(1) == "|":
            halves[n] = "a"
            n += "a"
        elif halves.get(n) == "a":
            halves[n] = "b"
            n += "b"
        out.append({"n": n, "text": text[start:m.start()].strip(" |\n"), "source": ""})
        start = m.end()
        src = SOURCE_RE.match(text, start)
        if src:
            out[-1]["source"] = src.group(1).strip()
            start = src.end()
    tail = text[start:].strip(" |\n")
    if tail:
        if out:
            out[-1]["text"] += " " + tail
        else:
            out.append({"n": "", "text": tail, "source": ""})
    return out


def _deva_num(n: str) -> str:
    return "".join(DEVA_DIGITS[int(c)] if c.isdigit() else c for c in n)


#: A commentary segment ends at a danda (| or ||); segments shorter than this
#: many characters are joined to the next, so a segment is about a sentence.
SEGMENT_MIN = 120
_DANDA_SPLIT = re.compile(r"(?<=\|)(?!\|)")


def segments(sa: str, iast: str, min_len: int = SEGMENT_MIN) -> List[Dict[str, Any]]:
    """A commentary cut into numbered segments at the edition's own sentence
    ends, the back-index between each English sentence and its Sanskrit.
    Each: {"k" (from 1), "sa", "iast"}. The IAST is cut at the same dandas
    and paired only when both scripts give the same number of pieces;
    otherwise each segment's iast is empty."""
    def pieces(t: str) -> List[str]:
        return [x.strip() for x in _DANDA_SPLIT.split(t.replace("\n", " ")) if x.strip()]
    ps, pi = pieces(sa), pieces(iast)
    paired = len(ps) == len(pi)
    out: List[Dict[str, Any]] = []
    cur_sa, cur_ia = "", ""
    for j, piece in enumerate(ps):
        cur_sa = (cur_sa + " " + piece).strip()
        if paired:
            cur_ia = (cur_ia + " " + pi[j]).strip()
        if len(cur_sa) >= min_len or j == len(ps) - 1:
            out.append({"k": len(out) + 1, "sa": cur_sa, "iast": cur_ia})
            cur_sa, cur_ia = "", ""
    return out


def work_file(edition: Dict[str, Any]) -> Dict[str, Any]:
    """What a translator is given: per section the heading, the root verses
    and each commentary as numbered segments. A translation answers each
    segment by its number (TRANSLATION_BRIEF)."""
    out = {"title_sa": edition["title_sa"], "title_iast": edition["title_iast"], "sections": []}
    for sec in edition["sections"]:
        out["sections"].append({
            "head_sa": sec["head_sa"], "head_iast": sec["head_iast"], "mula": sec["mula"],
            "madhukosa": segments(sec["madhukosa"]["sa"], sec["madhukosa"]["iast"]),
            "atanka": segments(sec["atanka"]["sa"], sec["atanka"]["iast"])})
    return out


def check_translation(edition: Dict[str, Any], tr_dir: Path) -> List[str]:
    """Every section has a file and every commentary segment an English
    answer, by number; returns what is missing or extra."""
    problems = []
    wf = work_file(edition)
    for i, sec in enumerate(wf["sections"]):
        path = tr_dir / f"sec_{i:02d}.json"
        if not path.is_file():
            problems.append(f"section {i}: no {path.name}")
            continue
        tr = json.loads(path.read_text(encoding="utf-8"))
        for key in ("madhukosa", "atanka"):
            if isinstance(tr.get(key + "_en"), str):
                # Chapters 1-2: the commentary translated whole, not by segment.
                if sec[key] and not tr[key + "_en"].strip():
                    problems.append(f"section {i} {key}: whole-text translation is empty")
                continue
            want = {s["k"] for s in sec[key]}
            got = {x.get("k") for x in tr.get(key) or []}
            if want - got:
                problems.append(f"section {i} {key}: segments {sorted(want - got)} untranslated")
            if got - want:
                problems.append(f"section {i} {key}: answers for segments {sorted(got - want)} that do not exist")
    return problems



def compose(edition: Dict[str, Any], tr_dir: Path, chapter: int, title_en: str, out: Path) -> Dict[str, Any]:
    import docx                                                 # python-docx
    d = docx.Document()
    d.add_heading(f"Chapter {chapter} — {edition['title_iast'].lstrip('0123456789. ')} "
                  f"({title_en}) · {edition['title_sa']}", level=1)
    report = {"sections": 0, "translated": 0, "verses": 0, "untranslated": []}

    halves_sa: Dict[str, str] = {}
    halves_iast: Dict[str, str] = {}

    def para(label: str, text: str) -> None:
        d.add_paragraph(f"{label}: {text}" if label else text)

    for i, sec in enumerate(edition["sections"]):
        report["sections"] += 1
        tr_path = tr_dir / f"sec_{i:02d}.json"
        tr = json.loads(tr_path.read_text(encoding="utf-8")) if tr_path.is_file() else None
        head = f"{tr['heading_en']} — " if tr else ""
        d.add_heading(f"{head}{sec['head_iast']} ({sec['head_sa']})", level=2)
        # The link back to the team's edition: its file and this section's
        # index in it (the edition's own sections, in order, from 0).
        para("Edition", f"{edition.get('file', '')} §{i}")
        if sec["head_sa"].strip() == COLOPHON:
            para("Note", f"Colophon: {sec['mula']['sa']} / {sec['mula']['iast']}")
        else:
            sa = split_verses(sec["mula"]["sa"], halves_sa)
            iast = split_verses(sec["mula"]["iast"], halves_iast)
            for v in sa:
                para("Verse", f"{v['text']} ॥ {_deva_num(v['n'])} ॥" if v["n"] else v["text"])
                report["verses"] += 1
            sources = sorted({v["source"] for v in sa if v["source"]})
            if sources:
                para("Source", "; ".join(sources))
            for v in iast:
                para("IAST", f"{v['text']} || {v['n']} ||" if v["n"] else v["text"])
        if tr is None:
            report["untranslated"].append(i)
            para("Note", "This section is not yet translated.")
        else:
            report["translated"] += 1
            if tr.get("padaccheda"):
                para("Padaccheda", tr["padaccheda"])
            if tr.get("words"):
                para("Words", "")
                for w in tr["words"]:
                    para("", f"{w['term']} — {w['en']}")
            para("Translation", tr.get("translation", ""))
        for key, label, en_key in (("madhukosa", "Madhukosha", "madhukosa_en"),
                                   ("atanka", "Atankadarpana", "atanka_en")):
            if not sec[key]["sa"]:
                continue
            if isinstance((tr or {}).get(key), list):
                # Segmented: [k] Sanskrit, then [k] English, per segment.
                en = {x["k"]: x.get("en", "") for x in tr[key]}
                for j, seg in enumerate(segments(sec[key]["sa"], sec[key]["iast"])):
                    para(label if j == 0 else "", f"[{seg['k']}] {seg['sa']}")
                    para("", f"[{seg['k']}] {en.get(seg['k'], '')}".rstrip())
                continue
            para(label, sec[key]["sa"].replace("\n", " "))
            for chunk in (tr or {}).get(en_key, "").split("\n\n"):
                if chunk.strip():
                    para("", chunk.strip())
        for n in (tr or {}).get("notes") or []:
            para("Note", n)
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(out)
    return report


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("edition", type=Path)
    ap.add_argument("translations", type=Path, nargs="?")
    ap.add_argument("out", type=Path, nargs="?")
    ap.add_argument("--chapter", type=int)
    ap.add_argument("--title-en")
    ap.add_argument("--work-file", type=Path,
                    help="write the translator's input (segmented) to this path and stop")
    ap.add_argument("--check", action="store_true",
                    help="report sections or segments the translation directory lacks, and stop")
    args = ap.parse_args()
    edition = read_edition(args.edition)
    if args.work_file:
        args.work_file.write_text(json.dumps(work_file(edition), ensure_ascii=False, indent=1),
                                  encoding="utf-8")
        return 0
    if args.check:
        problems = check_translation(edition, args.translations)
        print("\n".join(problems) or "complete")
        return 1 if problems else 0
    if not (args.translations and args.out and args.chapter and args.title_en):
        ap.error("composing needs translations, out, --chapter and --title-en")
    rep = compose(edition, args.translations, args.chapter, args.title_en, args.out)
    print(json.dumps(rep))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
